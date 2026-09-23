"""Adversarial authorization tests against an Alembic-migrated disposable PostgreSQL DB."""
import os
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import httpx
from openpyxl import load_workbook
import pytest
import pytest_asyncio
from sqlalchemy import select

from app.database import AsyncSessionLocal, engine
from app.main import app
from app.models.alert import Alert
from app.models.billing import BillingCustomer, Plan, Subscription
from app.models.device import Device
from app.models.organization import Organization, OrganizationMember
from app.models.project import UserProject
from app.models.relevance import FundingProject, OrganizationProfile
from app.models.saved_search import SavedSearch
from app.models.source import Source
from app.models.user import User
from app.models.workspace import DevicePipeline, FavoriteDevice
from app.utils.auth_utils import create_access_token
from app.services.match_service import _MATCH_STMT, _FALLBACK_MATCH_STMT
from app.services.notification_service import NotificationService
from app.services.billing_service import ensure_default_plans


pytestmark = [
    pytest.mark.asyncio(loop_scope="module"),
    pytest.mark.skipif(
        os.getenv("KAFUNDO_SECURITY_TESTS") != "1" or not os.getenv("TEST_POSTGRES_DATABASE_URL"),
        reason="Dedicated disposable PostgreSQL security database required",
    ),
]


def headers(user, *, claimed_role=None):
    token = create_access_token(str(user.id), claimed_role or user.role)
    return {"Authorization": f"Bearer {token}"}


@pytest_asyncio.fixture(loop_scope="module")
async def scenario():
    assert engine.url.database.startswith("kafundo_lot3_"), "Security tests require a disposable Lot 3 database"
    nonce = uuid4().hex[:12]
    users = {
        key: User(id=uuid4(), email=f"{key}-{nonce}@example.org", password_hash="unused", role="admin" if key == "global_admin" else "reader", full_name=key)
        for key in ("owner_a", "owner_b", "admin_a", "member_a", "viewer_a", "shared", "outsider", "global_admin", "unknown")
    }
    org_a = Organization(id=uuid4(), name=f"Org A {nonce}", slug=f"org-a-{nonce}", plan="free", status="active", created_by_id=users["owner_a"].id)
    org_b = Organization(id=uuid4(), name=f"Org B {nonce}", slug=f"org-b-{nonce}", plan="free", status="active", created_by_id=users["owner_b"].id)
    memberships = [
        ("owner_a", org_a, "org_owner"), ("owner_b", org_b, "org_owner"),
        ("admin_a", org_a, "org_admin"), ("member_a", org_a, "member"),
        ("viewer_a", org_a, "viewer"), ("shared", org_a, "member"),
        ("shared", org_b, "member"), ("unknown", org_a, "unexpected_role"),
    ]
    statuses = ("approved", "validated", "auto_published", "admin_only", "pending_review", "rejected")
    devices = {}
    for status in statuses:
        devices[status] = Device(
            id=uuid4(), title=f"Security {status} {nonce}", organism="Security issuer",
            country="France", device_type="subvention", status="open",
            source_url=f"https://example.org/{status}/{nonce}",
            short_description="A concrete funding opportunity for organizations in France. " * 3,
            full_description="Eligibility, funding and submission process are documented. " * 4,
            beneficiaries=["association"], funding_details="Funding is available for qualifying projects.",
            amount_max=10000, validation_status=status, user_quality_decision="publish",
            sectors=["energie"],
        )
    async with AsyncSessionLocal() as db:
        await ensure_default_plans(db)
        expert = (await db.execute(select(Plan).where(Plan.slug == "expert"))).scalar_one()
        expert.stripe_price_id = expert.stripe_price_id or f"price_expert_{nonce}"
        db.add_all(list(users.values()))
        await db.flush()
        db.add_all([org_a, org_b])
        await db.flush()
        for key, org, role in memberships:
            db.add(OrganizationMember(organization_id=org.id, user_id=users[key].id, role=role))
            if users[key].default_organization_id is None:
                users[key].default_organization_id = org.id
        db.add_all([
            Source(name=f"INTERNAL_SOURCE_{nonce}", organism="Test", country="France", source_type="portail_officiel", url=f"https://example.org/source/{nonce}", collection_mode="html", consecutive_errors=4),
            OrganizationProfile(organization_id=org_a.id, countries=["France"]),
            OrganizationProfile(organization_id=org_b.id, countries=["France"]),
            BillingCustomer(organization_id=org_a.id, stripe_customer_id=f"cus_a_{nonce}"),
            BillingCustomer(organization_id=org_b.id, stripe_customer_id=f"cus_b_{nonce}"),
        ])
        db.add_all(devices.values())
        await db.flush()
        pipeline = {}
        favorites = {}
        searches = {}
        alerts = {}
        user_projects = {}
        funding = {}
        for key, org, device in (
            ("owner_a", org_a, devices["approved"]),
            ("owner_b", org_b, devices["validated"]),
            ("shared_a", org_a, devices["auto_published"]),
            ("shared_b", org_b, devices["admin_only"]),
        ):
            user_key = "shared" if key.startswith("shared") else key
            user = users[user_key]
            pipeline[key] = DevicePipeline(
                user_id=user.id, organization_id=org.id, device_id=device.id,
                pipeline_status="soumis", priority="moyenne", note=f"SECRET_{key}_{nonce}",
                snapshot={"title": f"SECRET_{key}_{nonce}", "amountMax": 1000},
                documents=[{"id": f"doc-{key}", "name": f"DOC_{key}_{nonce}"}],
            )
            searches[key] = SavedSearch(user_id=user.id, organization_id=org.id, name=f"SEARCH_{key}_{nonce}", query={"q": key})
            favorites[key] = FavoriteDevice(user_id=user.id, organization_id=org.id, device_id=device.id, snapshot={"title": f"FAVORITE_{key}_{nonce}"})
            alerts[key] = Alert(user_id=user.id, organization_id=org.id, name=f"ALERT_{key}_{nonce}", criteria={})
            user_projects[key] = UserProject(user_id=user.id, organization_id=org.id, name=f"PROJECT_{key}_{nonce}", cached_matches=[{"id": str(devices["admin_only"].id), "title": f"PRIVATE_{nonce}"}])
            db.add_all([pipeline[key], favorites[key], searches[key], alerts[key], user_projects[key]])
        for key, org in (("a", org_a), ("b", org_b)):
            funding[key] = FundingProject(organization_id=org.id, created_by_id=users[f"owner_{key}"].id, name=f"FUNDING_{key}_{nonce}", status="active", is_primary=True)
            db.add(funding[key])
        await db.commit()
    return SimpleNamespace(
        users=users, a=org_a, b=org_b, devices=devices, pipeline=pipeline, favorites=favorites,
        searches=searches, alerts=alerts, user_projects=user_projects, funding=funding, nonce=nonce,
        paid_slug="expert",
    )


@pytest_asyncio.fixture(loop_scope="module")
async def client():
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as value:
        yield value
    await engine.dispose()


async def test_team_activity_reporting_and_shared_tenant_switch(scenario, client):
    shared = headers(scenario.users["shared"])
    team_a = await client.get("/api/v1/workspace/team", headers=shared)
    assert team_a.status_code == 200, team_a.text
    assert team_a.json()["organization_id"] == str(scenario.a.id)
    assert f"SECRET_shared_a_{scenario.nonce}" in team_a.text
    assert f"SECRET_shared_b_{scenario.nonce}" not in team_a.text
    assert f"SECRET_owner_b_{scenario.nonce}" not in team_a.text
    assert (await client.get("/api/v1/workspace/reporting", headers=shared)).json()["team_stats"]["total_tracked"] == 2
    assert f"SECRET_shared_b_{scenario.nonce}" not in (await client.get("/api/v1/workspace/activity", headers=shared)).text
    assert f"SEARCH_shared_b_{scenario.nonce}" not in (await client.get("/api/v1/workspace/saved-searches", headers=shared)).text
    assert f"FAVORITE_shared_b_{scenario.nonce}" not in (await client.get("/api/v1/workspace/favorites", headers=shared)).text
    assert (await client.post(f"/api/v1/organizations/{scenario.b.id}/select", headers=shared)).status_code == 200
    team_b = await client.get("/api/v1/workspace/team", headers=shared)
    assert team_b.json()["organization_id"] == str(scenario.b.id)
    assert f"SECRET_shared_b_{scenario.nonce}" in team_b.text
    assert f"SECRET_shared_a_{scenario.nonce}" not in team_b.text
    assert (await client.get("/api/v1/workspace/reporting", headers=shared)).json()["team_stats"]["total_tracked"] == 2
    assert f"SECRET_shared_a_{scenario.nonce}" not in (await client.get("/api/v1/workspace/activity", headers=shared)).text
    assert f"SEARCH_shared_a_{scenario.nonce}" not in (await client.get("/api/v1/workspace/saved-searches", headers=shared)).text
    assert f"FAVORITE_shared_a_{scenario.nonce}" not in (await client.get("/api/v1/workspace/favorites", headers=shared)).text
    assert (await client.post("/api/v1/workspace/pipeline", headers=shared, json={
        "device_id": str(scenario.devices["auto_published"].id), "pipeline_status": "soumis", "note": "overwrite",
    })).status_code == 404
    assert (await client.post("/api/v1/workspace/saved-searches", headers=shared, json={
        "id": str(scenario.searches["shared_a"].id), "name": "overwrite", "filters": {},
    })).status_code == 404
    assert (await client.post(f"/api/v1/organizations/{scenario.b.id}/select", headers=headers(scenario.users["owner_a"]))).status_code == 404
    assert (await client.get("/api/v1/workspace/team", headers=headers(scenario.users["outsider"]))).status_code == 403
    assert (await client.get("/api/v1/workspace/team", headers=headers(scenario.users["unknown"]))).status_code == 403
    reader_dashboard = await client.get("/api/v1/dashboard/", headers=headers(scenario.users["owner_a"]))
    assert reader_dashboard.status_code == 200, reader_dashboard.text
    assert f"INTERNAL_SOURCE_{scenario.nonce}" not in reader_dashboard.text
    internal_dashboard = await client.get("/api/v1/dashboard/", headers=headers(scenario.users["global_admin"]))
    assert internal_dashboard.json()["sources"]["in_error"] >= 1


async def test_idor_writes_and_viewer_denial(scenario, client):
    owner_a = headers(scenario.users["owner_a"])
    owner_b = headers(scenario.users["owner_b"])
    viewer = headers(scenario.users["viewer_a"])
    assert (await client.get(f"/api/v1/projects/{scenario.user_projects['owner_b'].id}", headers=owner_a)).status_code == 404
    assert (await client.put(f"/api/v1/projects/{scenario.user_projects['owner_b'].id}", headers=owner_a, json={"name": "stolen"})).status_code == 404
    assert (await client.delete(f"/api/v1/projects/{scenario.user_projects['owner_b'].id}", headers=owner_a)).status_code == 404
    assert (await client.get(f"/api/v1/alerts/{scenario.alerts['owner_b'].id}", headers=owner_a)).status_code == 404
    assert (await client.put(f"/api/v1/alerts/{scenario.alerts['owner_b'].id}", headers=owner_a, json={"name": "stolen"})).status_code == 404
    assert (await client.delete(f"/api/v1/alerts/{scenario.alerts['owner_b'].id}", headers=owner_a)).status_code == 404
    assert (await client.put(f"/api/v1/workspace/saved-searches/{scenario.searches['owner_b'].id}", headers=owner_a, json={"name": "stolen"})).status_code == 404
    assert (await client.delete(f"/api/v1/workspace/saved-searches/{scenario.searches['owner_b'].id}", headers=owner_a)).status_code == 204
    assert (await client.post(f"/api/v1/workspace/pipeline/{scenario.devices['validated'].id}/documents", headers=owner_a, json={"name": "stolen"})).status_code == 404
    assert (await client.put(f"/api/v1/funding-projects/{scenario.funding['b'].id}", headers=owner_a, json={"name": "stolen"})).status_code == 404
    assert (await client.delete(f"/api/v1/funding-projects/{scenario.funding['b'].id}", headers=owner_a)).status_code == 204
    assert (await client.get(f"/api/v1/alerts/{scenario.alerts['owner_a'].id}", headers=owner_b)).status_code == 404
    assert (await client.get("/api/v1/workspace/team", headers=viewer)).status_code == 200
    for method, url, payload in (
        ("post", "/api/v1/workspace/pipeline", {"device_id": str(scenario.devices["approved"].id), "pipeline_status": "soumis"}),
        ("post", "/api/v1/projects/", {"name": "Viewer project"}),
        ("put", "/api/v1/me/profile", {"organization_type": "association"}),
        ("post", "/api/v1/funding-projects", {"name": "Viewer funding"}),
        ("post", "/api/v1/alerts/", {"name": "Viewer alert"}),
    ):
        response = await getattr(client, method)(url, headers=viewer, json=payload)
        assert response.status_code == 403, (url, response.text)
    async with AsyncSessionLocal() as db:
        assert (await db.execute(select(Alert.name).where(Alert.id == scenario.alerts["owner_b"].id))).scalar_one() == f"ALERT_owner_b_{scenario.nonce}"
        assert (await db.execute(select(FundingProject.name).where(FundingProject.id == scenario.funding["b"].id))).scalar_one() == f"FUNDING_b_{scenario.nonce}"


async def test_publication_matrix_list_detail_search_export_and_matching(scenario, client):
    reader = headers(scenario.users["owner_a"])
    admin = headers(scenario.users["global_admin"])
    async with AsyncSessionLocal() as db:
        paid = (await db.execute(select(Plan).where(Plan.slug == scenario.paid_slug))).scalar_one()
        db.add(Subscription(organization_id=scenario.a.id, plan_id=paid.id, status="active"))
        await db.commit()
    public = {"approved", "validated", "auto_published"}
    for status, device in scenario.devices.items():
        expected = 200 if status in public else 404
        assert (await client.get(f"/api/v1/devices/{device.id}")).status_code == expected
        assert (await client.get(f"/api/v1/devices/{device.id}", headers=reader)).status_code == expected
        assert (await client.get(f"/api/v1/devices/{device.id}", headers=admin)).status_code == 200
    for request_headers in ({}, reader):
        listing = await client.get("/api/v1/devices/", params={"validation_status": "admin_only"}, headers=request_headers)
        assert listing.status_code == 200, listing.text
        assert all(str(item["id"]) != str(scenario.devices["admin_only"].id) for item in listing.json()["items"])
        searched = await client.get("/api/v1/devices/", params={"q": f"admin_only {scenario.nonce}"}, headers=request_headers)
        assert str(scenario.devices["admin_only"].id) not in searched.text
        exported = await client.get("/api/v1/devices/export/csv", params={"q": f"admin_only {scenario.nonce}"}, headers=request_headers)
        assert exported.status_code == (200 if request_headers else 401), exported.text
        if request_headers:
            assert f"Security admin_only {scenario.nonce}" not in exported.text
    assert str(scenario.devices["admin_only"].id) in (await client.get("/api/v1/devices/", params={"validation_status": "admin_only"}, headers=admin)).text
    assert (await client.get(f"/api/v1/devices/{scenario.devices['admin_only'].id}/relevance", headers=reader)).status_code == 404
    assert (await client.post(f"/api/v1/devices/{scenario.devices['admin_only'].id}/analyze", headers=reader)).status_code == 403
    assert (await client.post("/api/v1/projects/", headers=reader, json={
        "name": "Forged project", "organization_id": str(scenario.b.id), "user_id": str(scenario.users["owner_b"].id),
    })).status_code == 422
    before_refresh = await client.get(f"/api/v1/projects/{scenario.user_projects['owner_a'].id}", headers=reader)
    assert f"PRIVATE_{scenario.nonce}" not in before_refresh.text
    assert before_refresh.json()["match_score"] == 0
    assert (await client.post(f"/api/v1/projects/{scenario.user_projects['owner_a'].id}/match", headers=reader)).status_code == 200
    project = await client.get(f"/api/v1/projects/{scenario.user_projects['owner_a'].id}", headers=reader)
    assert f"PRIVATE_{scenario.nonce}" not in project.text
    assert f"Security admin_only {scenario.nonce}" not in project.text


async def test_billing_roles_and_mass_assignment(scenario, client, monkeypatch):
    calls = []

    def portal(**kwargs):
        calls.append(("portal", kwargs))
        return SimpleNamespace(url="https://stripe.example/portal")

    def checkout(**kwargs):
        calls.append(("checkout", kwargs))
        return SimpleNamespace(id="cs_test", url="https://stripe.example/checkout")

    fake_stripe = SimpleNamespace(
        billing_portal=SimpleNamespace(Session=SimpleNamespace(create=portal)),
        checkout=SimpleNamespace(Session=SimpleNamespace(create=checkout)),
        Subscription=SimpleNamespace(list=lambda **_kwargs: {"data": [], "has_more": False}),
    )
    monkeypatch.setattr("app.routers.billing._stripe", lambda: fake_stripe)
    roles = {key: headers(scenario.users[key]) for key in ("owner_a", "owner_b", "admin_a", "member_a", "viewer_a", "outsider", "shared")}
    for key in ("owner_a", "owner_b"):
        assert (await client.post("/api/v1/billing/portal", headers=roles[key])).status_code == 200
        assert (await client.post("/api/v1/billing/checkout", headers=roles[key], json={"plan_slug": scenario.paid_slug})).status_code == 200
    for key in ("admin_a", "member_a", "viewer_a", "outsider", "shared"):
        assert (await client.post("/api/v1/billing/portal", headers=roles[key])).status_code == 403
        assert (await client.post("/api/v1/billing/checkout", headers=roles[key], json={"plan_slug": scenario.paid_slug})).status_code == 403
    assert (await client.post("/api/v1/billing/portal")).status_code == 401
    assert (await client.post("/api/v1/billing/checkout", json={"plan_slug": scenario.paid_slug})).status_code == 401
    assert len(calls) == 4
    assert (await client.get("/api/v1/billing/subscription", headers=roles["viewer_a"])).status_code == 200
    assert (await client.get("/api/v1/billing/subscription", headers=roles["outsider"])).status_code == 403

    assert (await client.post("/api/v1/auth/register", json={
        "email": f"forged-{scenario.nonce}@example.org", "password": "123456789", "role": "admin",
    })).status_code == 422
    assert (await client.post("/api/v1/organizations", headers=roles["owner_a"], json={
        "name": "Forged org", "plan": "enterprise", "status": "active",
    })).status_code == 422
    assert (await client.post("/api/v1/funding-projects", headers=roles["owner_a"], json={
        "name": "Forged funding", "organization_id": str(scenario.b.id), "created_by_id": str(scenario.users["owner_b"].id),
    })).status_code == 422
    assert (await client.post("/api/v1/organizations/invite", headers=roles["owner_a"], json={
        "email": f"invite-{scenario.nonce}@example.org", "role": "member", "organization_id": str(scenario.b.id),
    })).status_code == 403
    forged_jwt = headers(scenario.users["owner_a"], claimed_role="admin")
    assert (await client.get(f"/api/v1/devices/{scenario.devices['admin_only'].id}", headers=forged_jwt)).status_code == 404


async def test_shared_billing_usage_and_invitation_roles(scenario, client, monkeypatch):
    shared = headers(scenario.users["shared"])
    for org in (scenario.a, scenario.b):
        assert (await client.post(f"/api/v1/organizations/{org.id}/select", headers=shared)).status_code == 200
        subscription = await client.get("/api/v1/billing/subscription", headers=shared)
        assert subscription.status_code == 200, subscription.text
        assert subscription.json()["organization_id"] == str(org.id)
        assert subscription.json()["usage"]["alerts"] == 2
        assert subscription.json()["usage"]["saved_searches"] == 2
        assert subscription.json()["usage"]["pipeline_projects"] == 2
    monkeypatch.setattr(NotificationService, "send_email", lambda *_args, **_kwargs: False)
    async with AsyncSessionLocal() as db:
        paid = (await db.execute(select(Plan).where(Plan.slug == scenario.paid_slug))).scalar_one()
        db.add(Subscription(organization_id=scenario.a.id, plan_id=paid.id, status="active"))
        await db.commit()
    for key in ("owner_a", "admin_a"):
        response = await client.post("/api/v1/organizations/invite", headers=headers(scenario.users[key]), json={
            "email": f"invite-{key}-{scenario.nonce}@example.org", "role": "member",
        })
        assert response.status_code == 201, response.text
        assert response.json()["organization_id"] == str(scenario.a.id)
    for key in ("member_a", "viewer_a", "outsider"):
        response = await client.post("/api/v1/organizations/invite", headers=headers(scenario.users[key]), json={
            "email": f"forbidden-{key}-{scenario.nonce}@example.org", "role": "member",
        })
        assert response.status_code in (403, 404), response.text
    assert (await client.post("/api/v1/organizations/invite", json={"email": "a@example.org", "role": "member"})).status_code == 401


async def test_match_sql_and_excel_export_exclude_unpublished(scenario, client):
    async with AsyncSessionLocal() as db:
        paid = (await db.execute(select(Plan).where(Plan.slug == scenario.paid_slug))).scalar_one()
        db.add(Subscription(organization_id=scenario.a.id, plan_id=paid.id, status="active"))
        await db.commit()
    params = {
        "query": "Security", "sectors": ["energie"], "countries": ["France"],
        "country_scopes": [], "regional_scopes": [], "geographic_scopes": [],
        "types": ["subvention"], "dominant_type": None,
        "amount_min": None, "amount_max": None, "limit": 50,
    }
    async with AsyncSessionLocal() as db:
        direct = (await db.execute(_MATCH_STMT, params)).mappings().all()
        fallback = (await db.execute(_FALLBACK_MATCH_STMT)).mappings().all()
    private_ids = {scenario.devices[status].id for status in ("admin_only", "pending_review", "rejected")}
    assert all(row["id"] not in private_ids for row in direct)
    assert all(row["id"] not in private_ids for row in fallback)
    for request_headers in ({}, headers(scenario.users["owner_a"])):
        response = await client.get("/api/v1/devices/export/excel", params={"q": f"admin_only {scenario.nonce}"}, headers=request_headers)
        assert response.status_code == (200 if request_headers else 401), response.text
        if not request_headers:
            continue
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        cells = " ".join(str(value) for row in workbook.active.iter_rows(values_only=True) for value in row if value is not None)
        assert f"Security admin_only {scenario.nonce}" not in cells
